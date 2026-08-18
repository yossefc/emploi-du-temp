"""Export d'un planning au FORMAT שחף — une feuille par classe, une par prof.

Reproduit la mise en page des exports de שחף (`emploi_du_temps_tashpau.xlsx`) :
ligne = période, colonne = jour, cellule = « matière\\nprof ». C'est le format
que l'école lit déjà, donc directement imprimable et comparable à l'existant.

Ce n'est PAS un import dans שחף : le logiciel exporte en Excel mais rien
n'indique qu'il sache relire un planning. Voir la discussion avec Yossef.

Usage :
    DATABASE_URL='sqlite:///./demo.db' python scripts/export_shahaf.py [schedule_id]
"""

from __future__ import annotations

import os
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("DATABASE_URL", "sqlite:///./demo.db")
os.environ.setdefault("SECRET_KEY", "this-is-a-test-secret-key-32chars-yes")

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import sessionmaker

from app.db.base import engine
import app.models  # noqa: F401
from app.models import Class, Group, Schedule, Subject, Teacher

OUT = Path(r"D:\emploi-du-temp-projet\donnees-tashpaz") / "מערכת_שעות_תשפז.xlsx"
DAYS_HE = ["ראשון", "שני", "שלישי", "רביעי", "חמישי"]
N_PERIODS = 10

THIN = Side(style="thin", color="AAB4C0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb, title: str, subtitle: str):
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = True
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    c = ws.cell(1, 1, subtitle)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1E3A5F")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.cell(2, 1, "שעה").font = Font(bold=True, color="FFFFFF")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="2E5C8A")
    ws.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(2, 1).border = BORDER
    for i, d in enumerate(DAYS_HE):
        c = ws.cell(2, i + 2, d)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5C8A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    ws.column_dimensions["A"].width = 6
    for col in "BCDEF":
        ws.column_dimensions[col].width = 26
    ws.row_dimensions[2].height = 20
    return ws


def _fill(ws, grid: dict, colors: dict):
    for p in range(N_PERIODS):
        r = p + 3
        c = ws.cell(r, 1, f"{p + 1}")
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        ws.row_dimensions[r].height = 34
        for d in range(5):
            items = grid.get((d, p), [])
            cell = ws.cell(r, d + 2, "\n".join(items))
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = BORDER
            cell.font = Font(size=9)
            key = colors.get((d, p))
            if key:
                cell.fill = PatternFill("solid", fgColor=key)


def main():
    db = sessionmaker(bind=engine)()
    sid = int(sys.argv[1]) if len(sys.argv) > 1 else None
    sched = (db.get(Schedule, sid) if sid else
             db.query(Schedule).order_by(Schedule.id.desc()).first())
    if sched is None:
        print("Aucun planning en base.")
        return
    school_id = sched.school_id
    groups = {g.id: g for g in db.query(Group).filter_by(school_id=school_id).all()}
    subjects = {s.id: s for s in db.query(Subject).filter_by(school_id=school_id).all()}
    classes = db.query(Class).filter_by(school_id=school_id).all()
    teachers = db.query(Teacher).filter_by(school_id=school_id).all()

    by_class: dict[int, dict] = defaultdict(lambda: defaultdict(list))
    by_teacher: dict[int, dict] = defaultdict(lambda: defaultdict(list))
    col_class: dict[int, dict] = defaultdict(dict)
    col_teacher: dict[int, dict] = defaultdict(dict)
    for en in sched.entries:
        g = groups[en.group_id]
        subj = subjects[g.subject_id]
        names = " / ".join(t.first_name for t in g.teachers)
        pos = (en.day_of_week, en.slot_index)
        hexcol = (subj.color_hex or "#94a3b8").lstrip("#").upper()
        for c in g.source_classes:
            by_class[c.id][pos].append(f"{subj.name_he}\n{names}")
            col_class[c.id][pos] = hexcol
        for t in g.teachers:
            label = ", ".join(sorted(x.code for x in g.source_classes)) or "ישיבה"
            by_teacher[t.id][pos].append(f"{subj.name_he}\n{label}")
            col_teacher[t.id][pos] = hexcol

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    order = ["ז", "ח", "ט", "י", "יא", "יב"]
    for cls in sorted(classes, key=lambda c: (order.index(c.code.split("-")[0]), c.code)):
        ws = _sheet(wb, cls.code, f'{cls.code} — מערכת שעות תשפ"ז')
        _fill(ws, by_class.get(cls.id, {}), col_class.get(cls.id, {}))
    n_t = 0
    for t in sorted(teachers, key=lambda t: t.first_name):
        if not by_teacher.get(t.id):
            continue
        ws = _sheet(wb, t.first_name, f'{t.first_name} — מערכת שעות תשפ"ז')
        _fill(ws, by_teacher[t.id], col_teacher.get(t.id, {}))
        n_t += 1

    wb.save(OUT)
    print(f"✅ {OUT}")
    print(f"   planning #{sched.id} « {sched.name} » — {len(sched.entries)} cours")
    print(f"   {len(classes)} feuilles classe + {n_t} feuilles prof")


if __name__ == "__main__":
    main()
