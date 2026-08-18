"""Deux classeurs récapitulatifs : matières × profs × heures.

Un fichier pour la חטיבה (ז/ח/ט), un pour le תיכון (י/יא/יב). Chacun comporte
trois vues de la même réalité :
  1. « לפי מקצוע »  — matière, prof, classes, heures ; le détail exhaustif.
  2. « לפי מורה »   — récapitulatif par enseignant, avec son total.
  3. « מטריצה »     — classes en colonnes, matières en lignes, heures dedans.

Une barrette (הקבצה) est signalée : elle occupe la classe le temps de son
enveloppe, pas la somme de ses groupes.

Usage : DATABASE_URL='sqlite:///./demo.db' python scripts/export_matieres.py
"""

from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("DATABASE_URL", "sqlite:///./demo.db")

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATA = Path(r"D:\emploi-du-temp-projet\donnees-tashpaz\dataset_tashpaz_final.json")
OUTDIR = Path(r"D:\emploi-du-temp-projet\donnees-tashpaz")

SEGMENTS = {
    "חטיבה": (("ז", "ח", "ט"), "מקצועות_ומורים_חטיבה.xlsx"),
    "תיכון": (("י", "יא", "יב"), "מקצועות_ומורים_תיכון.xlsx"),
}
GRADE_ORDER = ["ז", "ח", "ט", "י", "יא", "יב"]
THIN = Side(style="thin", color="BFC7D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="2E5C8A")
TITLE_FILL = PatternFill("solid", fgColor="1E3A5F")
ALT_FILL = PatternFill("solid", fgColor="F2F5F9")


def grade(code: str) -> str:
    return code.split("-")[0].strip()


def sheet(wb, title, subtitle, headers, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1, 1, subtitle)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    for j, h in enumerate(headers, 1):
        c = ws.cell(2, j, h)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[2].height = 30
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVWX", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    return ws


def write(ws, rows, start=3, numeric=()):
    for i, row in enumerate(rows):
        r = start + i
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v)
            c.border = BORDER
            c.alignment = Alignment(
                horizontal="center" if j in numeric else "right",
                vertical="center", wrap_text=(j == 3))
            if i % 2:
                c.fill = ALT_FILL
        ws.row_dimensions[r].height = 20
    return start + len(rows)


def build(seg_name, prefixes, out_name, data):
    groups = [g for g in data["groups"]
              if any(grade(c) in prefixes for c in g["classes"])]
    cohorts = data["cohorts"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- 1. par matière ----
    ws = sheet(wb, "לפי מקצוע", f'{seg_name} — מקצועות, מורים ושעות (תשפ"ז)',
               ["מקצוע", "מורה", "כיתות", "שעות", "הקבצה"],
               [20, 26, 30, 8, 34])
    rows = []
    for g in sorted(groups, key=lambda g: (g["subject"], -g["hours"], g["classes"][0])):
        coh = cohorts[g["cohort"]]["label"] if g["cohort"] is not None else ""
        who = g["teacher"] + ("  (לא משובץ)" if g["placeholder"] else "")
        rows.append([g["subject"], who, ", ".join(g["classes"]), g["hours"], coh])
    last = write(ws, rows, numeric=(4,))
    c = ws.cell(last, 1, 'סה"כ')
    c.font = Font(bold=True)
    c = ws.cell(last, 4, sum(g["hours"] for g in groups))
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")

    # ---- 2. par prof ----
    per = defaultdict(lambda: {"h": 0, "subj": set(), "cls": set(), "ph": False})
    for g in groups:
        p = per[g["teacher"]]
        p["h"] += g["hours"]
        p["subj"].add(g["subject"])
        p["cls"] |= set(g["classes"])
        p["ph"] = p["ph"] or g["placeholder"]
    ws = sheet(wb, "לפי מורה", f'{seg_name} — לפי מורה',
               ["מורה", "מקצועות", "כיתות", 'סה"כ שעות'], [26, 30, 34, 11])
    rows = [[n + ("  (לא משובץ)" if v["ph"] else ""),
             ", ".join(sorted(v["subj"])),
             ", ".join(sorted(v["cls"], key=lambda c: (GRADE_ORDER.index(grade(c)), c))),
             v["h"]]
            for n, v in sorted(per.items(), key=lambda x: -x[1]["h"])]
    last = write(ws, rows, numeric=(4,))
    ws.cell(last, 1, f'{len(per)} מורים').font = Font(bold=True)
    c = ws.cell(last, 4, sum(v["h"] for v in per.values()))
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")

    # ---- 3. matrice matières × classes ----
    cls_list = sorted({c for g in groups for c in g["classes"] if grade(c) in prefixes},
                      key=lambda c: (GRADE_ORDER.index(grade(c)), c))
    subj_list = sorted({g["subject"] for g in groups})
    # enveloppe : une barrette compte une fois par classe
    env = defaultdict(int)
    for g in data["groups"]:
        if g["cohort"] is not None:
            env[g["cohort"]] = max(env[g["cohort"]], g["hours"])
    grid = defaultdict(int)
    seen = defaultdict(set)
    for g in groups:
        for c in g["classes"]:
            if grade(c) not in prefixes:
                continue
            if g["cohort"] is not None:
                if g["cohort"] in seen[c]:
                    continue
                seen[c].add(g["cohort"])
                grid[(g["subject"], c)] += env[g["cohort"]]
            else:
                grid[(g["subject"], c)] += g["hours"]
    ws = sheet(wb, "מטריצה", f'{seg_name} — שעות לפי מקצוע וכיתה (הקבצה נספרת פעם אחת)',
               ["מקצוע"] + cls_list + ['סה"כ'],
               [20] + [8] * len(cls_list) + [9])
    rows = []
    for s in subj_list:
        vals = [grid.get((s, c), "") for c in cls_list]
        rows.append([s] + vals + [sum(v for v in vals if v)])
    rows.append(['סה"כ כיתה'] + [sum(grid.get((s, c), 0) for s in subj_list)
                                 for c in cls_list] + [""])
    write(ws, rows, numeric=tuple(range(2, len(cls_list) + 3)))
    for j in range(1, len(cls_list) + 3):
        ws.cell(2 + len(rows), j).font = Font(bold=True)

    out = OUTDIR / out_name
    wb.save(out)
    n_ph = sum(1 for v in per.values() if v["ph"])
    print(f"✅ {out.name}")
    print(f"   {len(groups)} cours · {len(subj_list)} matières · "
          f"{len(per)} profs (dont {n_ph} postes à pourvoir) · "
          f"{sum(g['hours'] for g in groups)}h")


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    for seg, (prefixes, name) in SEGMENTS.items():
        build(seg, prefixes, name, data)


if __name__ == "__main__":
    main()
