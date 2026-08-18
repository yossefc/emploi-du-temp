# -*- coding: utf-8 -*-
"""Export final : classes + profs + הקבצות. Args: <schedule_id|live> [out.xlsx]"""
from __future__ import annotations
import json, os, sys
from collections import defaultdict
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("DATABASE_URL", "sqlite:///./demo.db")
os.environ.setdefault("SECRET_KEY", "this-is-a-test-secret-key-32chars-yes")
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import sessionmaker
from app.db.base import engine
import app.models
from app.models import Group, ScheduleEntry

SCHOOL = 2
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("/home/claude/edt/מערכת_שעות_תשפז_final.xlsx")
DAYS = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי"]
db = sessionmaker(bind=engine)()
groups = {g.id: g for g in db.query(Group).filter_by(school_id=SCHOOL).all()}
slots_of = defaultdict(list)
src = sys.argv[1] if len(sys.argv) > 1 else "live"
if src.isdigit():
    for e in db.query(ScheduleEntry).filter_by(schedule_id=int(src)).all():
        slots_of[e.group_id].append((e.day_of_week, e.slot_index))
else:
    LIVE = Path(__file__).resolve().parents[2] / "warm_start" / f"school_{SCHOOL}_live.json"
    data = json.loads(LIVE.read_text(encoding="utf-8"))
    ref = data.get("groups", data)
    by_label = defaultdict(list)
    for g in groups.values(): by_label[g.label].append(g)
    used = set()
    for label, poss in ref.items():
        for g in by_label.get(label, []):
            if g.id in used: continue
            used.add(g.id); slots_of[g.id] = [tuple(x) for x in poss]; break
    print("meta:", data.get("_meta"))

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F4E5F")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
cell_al = Alignment(horizontal="center", vertical="center", wrap_text=True)
wb = Workbook(); wb.remove(wb.active)
def teachers_str(g): return " / ".join((t.first_name or "").strip() for t in g.teachers)

cls_cells = defaultdict(lambda: defaultdict(list))
for gid, poss in slots_of.items():
    g = groups[gid]; su = (g.subject.name_he or "").strip()
    if su == "ישיבות": continue
    for (d, s) in poss:
        for c in g.source_classes:
            cls_cells[c.code][(d, s)].append((su, teachers_str(g), g.parallel_cohort_id))
def cls_key(code):
    lvl, _, num = code.partition("-")
    return ({"ז":1,"ח":2,"ט":3,"י":4,"יא":5,"יב":6}.get(lvl.strip(), 9), num.strip())
ws = wb.create_sheet("לוח כיתות"); ws.sheet_view.rightToLeft = True
row = 1
for code in sorted(cls_cells, key=cls_key):
    cells = cls_cells[code]
    max_s = max(s for (_, s) in cells) + 1
    ws.cell(row, 1, f"כיתה {code}").font = Font(bold=True, size=14)
    row += 1; hr = row
    ws.cell(hr, 1, "שעה").fill = hdr_fill; ws.cell(hr, 1).font = hdr_font
    for d in range(6):
        c = ws.cell(hr, 2+d, DAYS[d]); c.fill = hdr_fill; c.font = hdr_font; c.alignment = cell_al
    for s in range(max_s):
        r = hr + 1 + s
        ws.cell(r, 1, f"P{s+1}").font = Font(bold=True)
        for d in range(6):
            items = cells.get((d, s), [])
            if not items:
                ws.cell(r, 2+d, "").border = border; continue
            subjects = sorted({it[0] for it in items})
            label = " + ".join(subjects) + ("  ⟨הקבצה⟩" if len(items) > 1 else "")
            txt = label + "\n" + " | ".join(it[1] for it in items)
            c = ws.cell(r, 2+d, txt); c.alignment = cell_al; c.border = border
            if len(items) > 1: c.fill = PatternFill("solid", fgColor="FFF3CD")
    row = hr + max_s + 2
for col in range(1, 8): ws.column_dimensions[get_column_letter(col)].width = 26 if col > 1 else 6

t_cells = defaultdict(lambda: defaultdict(list))
for gid, poss in slots_of.items():
    g = groups[gid]; su = (g.subject.name_he or "").strip()
    cls = "+".join(c.code for c in g.source_classes) or g.label
    for t in g.teachers:
        for (d, s) in poss: t_cells[(t.first_name or '').strip()][(d, s)].append(f"{su} {cls}")
ws = wb.create_sheet("לוח מורים"); ws.sheet_view.rightToLeft = True
row = 1
for nm in sorted(t_cells):
    if "לא משובץ" in nm: continue
    cells = t_cells[nm]; total = sum(len(v) for v in cells.values())
    ws.cell(row, 1, f"{nm}  ({total} ש')").font = Font(bold=True, size=13)
    row += 1; hr = row
    ws.cell(hr, 1, "שעה").fill = hdr_fill; ws.cell(hr, 1).font = hdr_font
    for d in range(6):
        c = ws.cell(hr, 2+d, DAYS[d]); c.fill = hdr_fill; c.font = hdr_font; c.alignment = cell_al
    max_s = max((s for (_, s) in cells), default=0) + 1
    for s in range(max_s):
        r = hr + 1 + s
        ws.cell(r, 1, f"P{s+1}").font = Font(bold=True)
        for d in range(6):
            c = ws.cell(r, 2+d, " | ".join(sorted(set(cells.get((d, s), [])))))
            c.alignment = cell_al; c.border = border
    row = hr + max_s + 2
for col in range(1, 8): ws.column_dimensions[get_column_letter(col)].width = 24 if col > 1 else 6

ws = wb.create_sheet("הקבצות"); ws.sheet_view.rightToLeft = True
for i, h in enumerate(["הקבצה","מקצוע","כיתות","מורה","שעות","משובץ","מתי"], 1):
    c = ws.cell(1, i, h); c.fill = hdr_fill; c.font = hdr_font
r = 2
cohorts = defaultdict(list)
for g in groups.values():
    if g.parallel_cohort_id: cohorts[g.parallel_cohort_id].append(g)
for cid in sorted(cohorts):
    for g in sorted(cohorts[cid], key=lambda g: -g.hours_per_week):
        poss = sorted(slots_of.get(g.id, []))
        when = ", ".join(f"{DAYS[d][:3]} P{s+1}" for d, s in poss)
        for i, v in enumerate([f"#{cid}", (g.subject.name_he or "").strip(),
                               "+".join(c.code for c in g.source_classes),
                               teachers_str(g), g.hours_per_week, len(poss), when], 1):
            ws.cell(r, i, v).border = border
        r += 1
    r += 1
for col, w in zip(range(1, 8), (8, 14, 18, 22, 7, 9, 60)):
    ws.column_dimensions[get_column_letter(col)].width = w
wb.save(OUT); print("OK →", OUT)
