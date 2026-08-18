"""Exporte en Excel la meilleure solution du solveur — PENDANT qu'il travaille.

Le moteur écrit une photo de sa meilleure solution chaque minute dans
warm_start/school_2_live.json. Ce script la transforme en classeur lisible,
sans toucher au calcul en cours.

    python scripts/voir_maintenant.py            # → מערכת_עכשיו.xlsx
    python scripts/voir_maintenant.py --check    # + contrôle des règles
"""
from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("DATABASE_URL", "sqlite:///./demo.db")
os.environ.setdefault("SECRET_KEY", "this-is-a-test-secret-key-32chars-yes")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import sessionmaker

from app.db.base import engine
import app.models  # noqa: F401
from app.models import Group

SCHOOL_ID = 2
LIVE = Path(__file__).resolve().parents[2] / "warm_start" / f"school_{SCHOOL_ID}_live.json"
OUT = Path(r"D:\emploi-du-temp-projet\donnees-tashpaz\מערכת_עכשיו.xlsx")
DAYS = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי"]
HATIVA = ("ז", "ח", "ט")
DISPERSED = {'חנ"ג', 'של"ח', "מנטורים", "חינוך", "ישיבות", "אמירים"}


def load():
    if not LIVE.exists():
        sys.exit(f"Pas encore de photo — le solveur n'a pas fini sa 1re solution.\n({LIVE})")
    data = json.loads(LIVE.read_text(encoding="utf-8"))
    return data.get("_meta", {}), data.get("groups", data)


def main() -> None:
    meta, ref = load()
    db = sessionmaker(bind=engine)()
    groups = db.query(Group).filter_by(school_id=SCHOOL_ID).all()
    by_label = defaultdict(list)
    for g in groups:
        by_label[g.label].append(g)

    grid: dict[str, dict[tuple[int, int], tuple[str, str, str]]] = defaultdict(dict)
    subj_day: dict[tuple[str, str, int], set[int]] = defaultdict(set)
    for label, poss in ref.items():
        gs = by_label.get(label)
        if not gs:
            continue
        g = gs[0]
        su = (g.subject.name_he or g.subject.name_fr or "").strip()
        te = " / ".join(t.first_name for t in g.teachers)
        col = (g.subject.color_hex or "#94A3B8").lstrip("#")
        for d, s in poss:
            for c in g.source_classes:
                prev = grid[c.code].get((d, s))
                if prev and su in prev[0]:
                    continue
                grid[c.code][(d, s)] = (
                    f"{prev[0]} / {su}" if prev else su,
                    f"{prev[1]} / {te}" if prev else te, col)
                subj_day[(c.code, su, d)].add(s)

    print(f"Photo : pénalité {meta.get('penalty', '?'):,} · "
          f"{meta.get('minutes', '?')} min · phase {meta.get('phase', '?')}")

    if "--check" in sys.argv:
        gaps = late = triple = 0
        for code, cells in grid.items():
            byd = defaultdict(list)
            for (d, s) in cells:
                if d < 5:
                    byd[d].append(s)
            for d, sl in byd.items():
                sl = sorted(sl)
                gaps += (sl[-1] - sl[0] + 1) - len(sl)
                late += sl[0] != 0
        for (code, su, d), sl in subj_day.items():
            if code.split("-")[0] not in HATIVA or su in DISPERSED:
                continue
            sl = sorted(sl)
            run = mx = 1
            for i in range(1, len(sl)):
                run = run + 1 if sl[i] == sl[i - 1] + 1 else 1
                mx = max(mx, run)
            triple += mx >= 3
        print(f"  trous élèves {gaps} · départs hors P1 {late} · חטיבה 3h+ {triple}")

    thin = Side(style="thin", color="CBD5E1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    order = {"ז": 0, "ח": 1, "ט": 2, "י": 3, "יא": 4, "יב": 5}

    def key(c: str):
        p = c.split("-")
        return (order.get(p[0], 9), int(p[1]) if len(p) > 1 and p[1].isdigit() else 0)

    wb = Workbook()
    wb.remove(wb.active)
    for code in sorted(grid, key=key):
        ws = wb.create_sheet(code)
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions["A"].width = 7
        for i in range(2, 8):
            ws.column_dimensions[get_column_letter(i)].width = 21
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
        t = ws.cell(1, 1, f'כיתה {code} — תמונת מצב ({meta.get("minutes", "?")} דק\')')
        t.font = Font(bold=True, size=15, color="FFFFFF", name="Arial")
        t.fill = PatternFill("solid", fgColor="1E3A5F")
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34
        for j, d in enumerate(DAYS):
            h = ws.cell(2, j + 2, d)
            h.font = Font(bold=True, size=12, color="FFFFFF", name="Arial")
            h.fill = PatternFill("solid", fgColor="2E5F8A")
            h.alignment = Alignment(horizontal="center", vertical="center")
            h.border = box
        ws.row_dimensions[2].height = 26
        last = max((s for (_, s) in grid[code]), default=9)
        for s in range(last + 1):
            r = s + 3
            p = ws.cell(r, 1, f"P{s+1}")
            p.font = Font(bold=True, size=11, color="1E3A5F", name="Arial")
            p.fill = PatternFill("solid", fgColor="E2E8F0")
            p.alignment = Alignment(horizontal="center", vertical="center")
            p.border = box
            ws.row_dimensions[r].height = 38
            for j in range(6):
                cell = ws.cell(r, j + 2)
                cell.border = box
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True, readingOrder=2)
                v = grid[code].get((j, s))
                if v:
                    cell.value = f"{v[0]}\n{v[1]}"
                    cell.font = Font(size=10, bold=True, color="FFFFFF", name="Arial")
                    cell.fill = PatternFill("solid", fgColor=v[2])
    wb.save(OUT)
    print(f"✅ {OUT} — {len(grid)} classes")


if __name__ == "__main__":
    main()
